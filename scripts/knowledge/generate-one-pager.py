#!/usr/bin/env python3
"""Build the published Knowledge concept one-pagers as Korean PDF and DOCX files."""

import json
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[2]
REGISTRY = ROOT / "content/knowledge/registry.json"
CONTENT = ROOT / "content/knowledge"
OUTPUT = ROOT / "public/knowledge"
FONT = "NotoSansKR"
FONT_PATH = Path(__file__).resolve().parent / "assets/NotoSansKR-Hangul.ttf"
INK = "#111820"
MUTED = "#48545E"
FLOW_COLORS = ["#F0E0E3", "#DCE6EC", "#D9E6DF", "#EEE7D7", "#E4DFF0"]


def normalized_one_pager(article, summary):
    raw = article["onePager"]
    title = raw.get("title") or article["canonicalName"]["ko"]
    subtitle = raw.get("subtitle") or article["dek"]["ko"]
    one_line = raw.get("oneLine") or raw.get("oneLineDefinition") or raw.get("sections", [{}])[0].get("body", "")
    sources = raw.get("sources") or [item["url"] for item in article.get("sources", [])]
    flow = raw.get("flow") or ["전경 레이어", "Alpha 계산", "배경 결합", "합성 화면"]
    paths = summary.get("onePager", {})
    pdf_path = raw.get("pdfPath") or paths.get("pdfPath")
    docx_path = raw.get("docxPath") or paths.get("docxPath")
    if not pdf_path or not docx_path:
        raise ValueError(f"Missing PDF/DOCX paths for {article['slug']}.")
    return {
        "title": title,
        "subtitle": subtitle,
        "oneLine": one_line,
        "sections": raw["sections"],
        "sources": sources,
        "flow": flow,
        "pdf": OUTPUT / Path(pdf_path).name,
        "docx": OUTPUT / Path(docx_path).name,
    }


def flow_table(flow, styles, width):
    steps = flow[:5]
    if len(steps) < 2:
        return None
    arrow_width = 17
    step_width = (width - arrow_width * (len(steps) - 1)) / len(steps)
    cells = []
    widths = []
    for index, step in enumerate(steps):
        cells.append(Paragraph(escape(str(step)), styles["flow"]))
        widths.append(step_width)
        if index < len(steps) - 1:
            cells.append(Paragraph("→", styles["arrow"]))
            widths.append(arrow_width)
    table = Table([cells], colWidths=widths, rowHeights=None, hAlign="LEFT")
    commands = [
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor(INK)),
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("BOX", (0, 0), (-1, -1), 0.45, colors.HexColor("#C7D0D6")),
    ]
    for index in range(0, len(cells), 2):
        commands.append(("BACKGROUND", (index, 0), (index, 0), colors.HexColor(FLOW_COLORS[(index // 2) % len(FLOW_COLORS)])))
    table.setStyle(TableStyle(commands))
    return table


def build_pdf(item):
    if FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(FONT, str(FONT_PATH), subfontIndex=0))
    styles = {
        "eyebrow": ParagraphStyle("eyebrow", fontName=FONT, fontSize=7.2, leading=9, textColor=colors.HexColor("#637783"), spaceAfter=5, wordWrap="CJK"),
        "title": ParagraphStyle("title", fontName=FONT, fontSize=18, leading=22, textColor=colors.HexColor(INK), spaceAfter=3, wordWrap="CJK"),
        "subtitle": ParagraphStyle("subtitle", fontName=FONT, fontSize=9, leading=12, textColor=colors.HexColor(MUTED), spaceAfter=8, wordWrap="CJK"),
        "definition": ParagraphStyle("definition", fontName=FONT, fontSize=9.3, leading=13.5, textColor=colors.HexColor(INK), spaceAfter=7, wordWrap="CJK"),
        "heading": ParagraphStyle("heading", fontName=FONT, fontSize=8.8, leading=11, textColor=colors.HexColor(INK), spaceBefore=2, spaceAfter=1, wordWrap="CJK"),
        "body": ParagraphStyle("body", fontName=FONT, fontSize=7.9, leading=11.2, textColor=colors.HexColor(MUTED), spaceAfter=3, wordWrap="CJK", alignment=TA_LEFT),
        "source": ParagraphStyle("source", fontName=FONT, fontSize=6.2, leading=7.5, textColor=colors.HexColor(MUTED), spaceAfter=1, wordWrap="CJK"),
        "flow": ParagraphStyle("flow", fontName=FONT, fontSize=7, leading=9, textColor=colors.HexColor(INK), alignment=1, wordWrap="CJK"),
        "arrow": ParagraphStyle("arrow", fontName=FONT, fontSize=8, leading=9, textColor=colors.HexColor("#637783"), alignment=1),
    }
    story = [
        Paragraph("RESEARCH KNOWLEDGE / CONCEPT ONE-PAGER", styles["eyebrow"]),
        Paragraph(escape(item["title"]), styles["title"]),
        Paragraph(escape(item["subtitle"]), styles["subtitle"]),
        Paragraph(f"<b>01. 한 줄 정의</b>  {escape(item['oneLine'])}", styles["definition"]),
    ]
    diagram = flow_table(item["flow"], styles, A4[0] - 68)
    if diagram:
        story.extend([diagram, Spacer(1, 5)])

    for index, section in enumerate(item["sections"], start=2):
        story.append(KeepTogether([
            Paragraph(f"{index:02d}. {escape(section['heading'])}", styles["heading"]),
            Paragraph(escape(section["body"]), styles["body"]),
        ]))
    story.extend([Spacer(1, 2), Paragraph("SOURCES", styles["eyebrow"])])
    for source in item["sources"]:
        story.append(Paragraph(escape(str(source)), styles["source"]))

    document = SimpleDocTemplate(
        str(item["pdf"]), pagesize=A4, leftMargin=34, rightMargin=34,
        topMargin=27, bottomMargin=25, title=item["title"], author="Minho Lee",
    )
    document.build(story)


def set_cell_shading(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shading)


def set_table_borders(table, color="D2D9DE"):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    table._tbl.tblPr.append(borders)


def set_run_font(run, size, bold=False, color=INK):
    run.font.name = "Noto Sans KR"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans KR")


def add_text(doc, text, size=8.1, bold=False, color=MUTED, before=0, after=2, line=1.0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    set_run_font(paragraph.add_run(str(text)), size, bold, color)
    return paragraph


def add_flow_to_docx(doc, flow):
    steps = flow[:5]
    if len(steps) < 2:
        return
    labels = []
    for index, step in enumerate(steps):
        labels.append(str(step))
        if index < len(steps) - 1:
            labels.append("→")
    table = doc.add_table(rows=1, cols=len(labels))
    table.autofit = False
    total_width = 18.4
    arrow_count = len(steps) - 1
    step_width = (total_width - arrow_count * 0.65) / len(steps)
    for index, (cell, label) in enumerate(zip(table.rows[0].cells, labels)):
        arrow = index % 2 == 1
        cell.width = Cm(0.65 if arrow else step_width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if not arrow:
            set_cell_shading(cell, FLOW_COLORS[index // 2 % len(FLOW_COLORS)])
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 0.95
        set_run_font(paragraph.add_run(label), 6.5 if arrow else 6.8, arrow, INK)
    set_table_borders(table)
    add_text(doc, "", 1, after=2)


def build_docx(item):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.95)
    section.bottom_margin = Cm(0.9)
    section.left_margin = Cm(1.15)
    section.right_margin = Cm(1.15)
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans KR"
    normal.font.size = Pt(8.1)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans KR")
    normal.paragraph_format.space_after = Pt(2)

    add_text(doc, "RESEARCH KNOWLEDGE / CONCEPT ONE-PAGER", 7.2, True, "#637783", after=3)
    add_text(doc, item["title"], 17, False, INK, after=1, line=1.0)
    add_text(doc, item["subtitle"], 8.8, False, MUTED, after=5)
    add_text(doc, "01. 한 줄 정의", 8.8, True, INK, after=1)
    add_text(doc, item["oneLine"], 8.1, False, MUTED, after=4)
    add_flow_to_docx(doc, item["flow"])

    for index, entry in enumerate(item["sections"], start=2):
        add_text(doc, f"{index:02d}. {entry['heading']}", 8.5, True, INK, before=1, after=0.5)
        add_text(doc, entry["body"], 7.8, False, MUTED, after=2, line=1.0)
    add_text(doc, "SOURCES", 7.2, True, "#637783", before=2, after=1)
    for source in item["sources"]:
        add_text(doc, source, 6, False, MUTED, after=0, line=1.0)
    doc.core_properties.title = item["title"]
    doc.core_properties.subject = "A4 Concept One-Pager"
    doc.core_properties.author = "Minho Lee"
    doc.save(item["docx"])


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    registry = json.loads(REGISTRY.read_text(encoding="utf-8"))
    generated = []
    for summary in registry["concepts"]:
        if summary["status"] != "published":
            continue
        article = json.loads((CONTENT / f"{summary['slug']}.json").read_text(encoding="utf-8"))
        item = normalized_one_pager(article, summary)
        build_pdf(item)
        build_docx(item)
        generated.append(summary["slug"])
    print(f"Created {len(generated)} Knowledge One-Pager PDF/DOCX pairs: {', '.join(generated)}")


if __name__ == "__main__":
    main()
